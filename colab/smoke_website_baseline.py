"""Smoke test for the composition baseline and Excel columns on branch fix/website-claims.

30 small molecules, two RDKit-computed targets, no network. Run from the repository root:
    python colab/smoke_website_baseline.py
"""
import sys, os, tempfile, math
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import pandas as pd
from rdkit import Chem
from rdkit.Chem import Descriptors
import ml, main
smis = ["CCO","CCCO","CCCCO","CC(C)O","c1ccccc1","Cc1ccccc1","CCc1ccccc1","c1ccc(O)cc1","CC(=O)O","CCC(=O)O",
        "CCCCC(=O)O","c1ccncc1","Cc1ccncc1","CCN","CCCN","CCCCN","ClCCCl","BrCCBr","CC(C)(C)O","C1CCCCC1",
        "C1CCCC1","CC1CCCCC1","OCCO","OCC(O)CO","CC(=O)OC","CCOC(=O)C","c1ccc2ccccc2c1","CC(N)C(=O)O","NCC(=O)O","CSC"]
df = pd.DataFrame({"SMILES": smis})
df["MolWt"] = [Descriptors.MolWt(Chem.MolFromSmiles(s)) for s in smis]
df["LogP"] = [Descriptors.MolLogP(Chem.MolFromSmiles(s)) for s in smis]
df = ml.compute_topological_indices(df)
res, best = ml.run_ml_qspr(df)
top = res.sort_values("R2_mean", ascending=False).drop_duplicates("Property")
print(top[["Property","Model","R2_mean","Comp_R2_mean","Delta_R2"]].to_string(index=False))
assert {"Comp_R2_mean","Delta_R2"} <= set(res.columns)
with tempfile.TemporaryDirectory() as d:
    out = os.path.join(d, "x.xlsx")
    main.export_excel("smoke", df, ml.run_correlation(df), res, top.reset_index(drop=True), {}, out)
    import openpyxl
    ws = openpyxl.load_workbook(out)["Best Models"]
    print([c.value for c in ws[1]]); print([c.value for c in ws[2]])
comp = ml._composition_matrix(["CCO"]).tolist()
print("composition of CCO (C,H,N,O,S,P,F,Cl,Br,I,other):", comp)
assert comp == [[2, 6, 0, 1, 0, 0, 0, 0, 0, 0, 0]]
assert list(top.Delta_R2.round(4)) == list((top.R2_mean - top.Comp_R2_mean).round(4))
print("SMOKE TEST PASSED")
